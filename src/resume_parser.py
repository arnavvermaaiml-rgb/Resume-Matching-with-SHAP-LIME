"""
Resume parsing: turn an uploaded PDF / DOCX / TXT resume into the same
structured record the model was trained on (skills, technologies, years of
experience, education level + field, projects, full text).

This is a rule-based parser (regex + vocabularies), not an LLM, so it is
fast, free and deterministic -- but it will miss unusual layouts. Every parse
returns `warnings` describing anything it could not find, and the app shows them.
"""
import io
import re
from datetime import date
from pathlib import Path

from skills import ALIASES, mentions, normalize_skill

# ---------------------------------------------------------------------------
# Vocabularies
# ---------------------------------------------------------------------------
BUILTIN_SKILLS = sorted({
    # languages
    "python", "java", "javascript", "typescript", "c++", "c#", "golang", "rust", "kotlin", "swift",
    "php", "ruby", "scala", "matlab", "bash", "sql", "html", "css",
    # data / ML
    "machine learning", "deep learning", "nlp", "computer vision", "pandas", "numpy", "scikit-learn",
    "tensorflow", "pytorch", "keras", "statistics", "tableau", "power bi", "excel", "nosql", "mongodb",
    "postgresql", "mysql", "redis", "spark", "hadoop", "airflow", "kafka", "etl", "data pipelines",
    "data analysis", "data visualization", "mlops", "llm", "feature engineering", "data modeling",
    # web / backend
    "react", "angular", "vue", "node.js", "next.js", "redux", "rest apis", "graphql", "django", "flask",
    "fastapi", "spring boot", "microservices", "webpack", "testing",
    # devops / cloud
    "docker", "kubernetes", "terraform", "ansible", "jenkins", "ci/cd", "linux", "git", "aws", "azure",
    "gcp", "cloud architecture", "networking", "security", "monitoring", "prometheus", "grafana",
    "cost optimization",
    # QA
    "selenium", "cypress", "pytest", "test automation", "api testing", "performance testing", "junit",
    # product / design
    "product strategy", "roadmapping", "stakeholder management", "agile", "scrum", "jira", "a/b testing",
    "ux research", "figma", "adobe xd", "wireframing", "prototyping", "ux design", "user research",
    "usability testing", "design systems", "analytics",
})

# Short or everyday-English aliases that would create false positives when scanning free text.
AMBIGUOUS_ALIASES = {"spring", "torch", "dl", "ts", "stats"}

ALIAS_VARIANTS = {}
for _alias, _canon in ALIASES.items():
    if _alias not in AMBIGUOUS_ALIASES:
        ALIAS_VARIANTS.setdefault(_canon, []).append(_alias)

KNOWN_FIELDS = [
    "Computer Science", "Software Engineering", "Information Technology", "Computer Engineering",
    "Artificial Intelligence", "Data Science", "Statistics", "Mathematics", "Electrical Engineering",
    "Mechanical Engineering", "Physics", "Business Analytics", "Business Administration", "Economics",
    "Human-Computer Interaction", "Graphic Design", "Cognitive Science", "Psychology", "Marketing",
    "Fine Arts", "History", "Hospitality Management", "Design", "Engineering",
]

DEGREE_PATTERNS = [  # (level, regex) -- checked highest level first
    ("phd", r"\bph\.?\s?d\b|\bdoctorate\b|\bdoctor of philosophy\b"),
    ("master", r"\bmaster'?s?\b|\bm\.?\s?sc\b|\bm\.?\s?tech\b|\bm\.?\s?s\b(?=[\s.,]|$)|\bmba\b|\bm\.?\s?eng\b|\bm\.?\s?e\b(?=[\s.,]|$)"),
    ("bachelor", r"\bbachelor'?s?\b|\bb\.?\s?sc\b|\bb\.?\s?tech\b|\bb\.?\s?s\b(?=[\s.,]|$)|\bb\.?\s?a\b(?=[\s.,]|$)|\bb\.?\s?e\b(?=[\s.,]|$)|\bundergraduate\b|\bb\.?\s?eng\b"),
    ("associate", r"\bassociate'?s?\b|\bdiploma\b"),
]
EDU_LABEL = {"associate": "Associate's", "bachelor": "Bachelor's", "master": "Master's", "phd": "PhD"}

SECTION_HEADINGS = {
    "skills": r"(?:technical\s+|key\s+|core\s+)?(?:skills|competencies|technologies|tech stack|skills\s*(?:&|and)\s*tools)",
    "experience": r"(?:work\s+|professional\s+|relevant\s+)?(?:experience|employment(?:\s+history)?|work history)",
    "education": r"education(?:al)?(?:\s+background)?|academics|qualifications",
    "projects": r"(?:personal\s+|academic\s+|key\s+|selected\s+)?projects",
    "summary": r"(?:professional\s+)?(?:summary|profile|objective|about(?:\s+me)?)",
    "other": r"certifications?|achievements?|awards?|publications?|interests|hobbies|languages|references|extracurriculars?|volunteer(?:ing)?",
}

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], start=1)}


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------
def extract_text(data: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if ext == ".docx":
        import docx
        document = docx.Document(io.BytesIO(data))
        parts = [("- " + p.text) if p.style.name.lower().startswith("list") and p.text.strip() else p.text
                 for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append("  ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    if ext in (".txt", ".md", ""):
        for enc in ("utf-8", "latin-1"):
            try:
                return data.decode(enc)
            except UnicodeDecodeError:
                continue
    raise ValueError(f"Unsupported file type '{ext}'. Upload a PDF, DOCX or TXT resume.")


def _clean(text: str) -> str:
    text = text.replace("•", "\n- ").replace("●", "\n- ").replace("▪", "\n- ").replace("", "\n- ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


# ---------------------------------------------------------------------------
# Section splitting
# ---------------------------------------------------------------------------
def split_sections(text: str) -> dict:
    sections = {"header": []}
    current = "header"
    for line in text.splitlines():
        stripped = line.strip().rstrip(":").strip()
        matched = None
        if stripped and len(stripped) <= 45:
            for name, pattern in SECTION_HEADINGS.items():
                if re.fullmatch(pattern, stripped, flags=re.I):
                    matched = name
                    break
        if matched:
            current = matched
            sections.setdefault(current, [])
        else:
            sections.setdefault(current, []).append(line)
    return {k: "\n".join(v).strip() for k, v in sections.items()}


# ---------------------------------------------------------------------------
# Field extractors
# ---------------------------------------------------------------------------
def extract_name(text: str, filename: str = "") -> str:
    first = True
    for line in text.splitlines()[:8]:
        line = line.strip()
        if not line:
            continue
        if any(re.fullmatch(p, line.rstrip(":"), flags=re.I) for p in SECTION_HEADINGS.values()):
            break  # reached the first section: the name would have come before it
        is_first, first = first, False
        if re.search(r"[\d@/|:]", line) or len(line) > 40:
            continue
        words = line.split()
        # a name is capitalised words ("and decisions." from a wrapped summary line is not), and a lone word
        # only counts when it is the very first line of the resume
        if (2 <= len(words) <= 4 or (is_first and len(words) == 1)) and all(
            re.fullmatch(r"[A-Z][A-Za-z'\-]*\.?", w) and (not w.endswith(".") or len(w) == 2) for w in words
        ):
            return line.title() if line.isupper() else line
    return Path(filename).stem.replace("_", " ").replace("-", " ").title() or "Unknown candidate"


def extract_email(text: str) -> str:
    m = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)
    return m.group(0) if m else ""


def detect_skills(text: str, extra_terms=()) -> list:
    """Canonical skills found in the text: built-in vocabulary plus any job-specific terms."""
    vocab = {normalize_skill(t) for t in BUILTIN_SKILLS} | {normalize_skill(t) for t in extra_terms}
    found = []
    text = re.sub(r"\s+", " ", text)  # PDFs wrap lines mid-skill ("A/B\nTesting")
    scan = text
    # Multi-word skills first, masked once found, so "API testing" doesn't also count as plain "testing".
    for canon in sorted(vocab, key=lambda c: (-len(c.split()), c)):
        variants = [canon] + ALIAS_VARIANTS.get(canon, [])
        hit = False
        for v in variants:
            if mentions(scan, v):
                hit = True
                if len(v.split()) > 1:
                    scan = re.sub(r"(?<![a-z0-9])" + re.escape(v) + r"(?![a-z0-9])", " ", scan, flags=re.I)
        if hit:
            found.append(canon)
    # "Spring" alone is ambiguous (the season); treat it as Spring Boot only alongside Java.
    if "spring boot" not in found and mentions(text, "spring") and mentions(text, "java"):
        found.append("spring boot")
    return sorted(found)


def _to_month_index(token: str, today: date):
    token = token.strip().lower()
    if token in ("present", "current", "now", "till date", "to date", "ongoing"):
        return today.year * 12 + today.month
    m = re.fullmatch(r"([a-z]{3,9})\.?,?\s*(\d{4})", token)
    if m and m.group(1)[:3] in MONTHS:
        return int(m.group(2)) * 12 + MONTHS[m.group(1)[:3]]
    m = re.fullmatch(r"(\d{1,2})[/.-](\d{4})", token)
    if m and 1 <= int(m.group(1)) <= 12:
        return int(m.group(2)) * 12 + int(m.group(1))
    m = re.fullmatch(r"(\d{4})", token)
    if m and 1950 <= int(m.group(1)) <= today.year + 1:
        return int(m.group(1)) * 12 + 1
    return None


def years_from_date_ranges(text: str, today: date = None) -> float:
    """Total months covered by date ranges (overlaps merged), in years."""
    today = today or date.today()
    month = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*"
    tok = rf"(?:{month}\.?,?[ \t]*\d{{4}}|\d{{1,2}}[/.-]\d{{4}}|\d{{4}}|present|current|now|ongoing|till date|to date)"
    pattern = re.compile(rf"({tok})\s*(?:-|–|—|to|until)\s*({tok})", re.I)
    intervals = []
    for a, b in pattern.findall(text):
        start, end = _to_month_index(a, today), _to_month_index(b, today)
        if start is None or end is None or end < start:
            continue
        # bare "2019 - 2021" ranges mean whole calendar years
        if re.fullmatch(r"\d{4}", b.strip()):
            end += 11
        intervals.append((start, end))
    if not intervals:
        return 0.0
    intervals.sort()
    merged = [list(intervals[0])]
    for s, e in intervals[1:]:
        if s <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], e)
        else:
            merged.append([s, e])
    return sum(e - s for s, e in merged) / 12.0


def years_stated(text: str) -> float:
    best = 0.0
    patterns = [
        r"(\d{1,2})(?:\.\d)?\s*\+?\s*(?:years?|yrs?)\s+(?:of\s+)?(?:\w+[\s,/&-]+){0,4}?experience",
        r"experience\s*(?:of|:)?\s*(\d{1,2})\s*\+?\s*(?:years?|yrs?)",
        r"(\d{1,2})\s*\+?\s*(?:years?|yrs?)\s+(?:in the field|of (?:professional|industry|work))",
    ]
    for p in patterns:
        for m in re.finditer(p, text, flags=re.I):
            best = max(best, float(m.group(1)))
    if best == 0.0:  # looser fallback: "spent 5 years building ..."
        for m in re.finditer(r"(\d{1,2})\s*\+?\s*(?:years?|yrs?)\b", text, flags=re.I):
            after = text[m.end(): m.end() + 25].lower()  # "4 years degree/program" is schooling, not work
            if not re.match(r"\s*(?:of\s+)?(?:degree|bachelor|master|course|program|school|college|university)", after):
                best = max(best, float(m.group(1)))
    return best


def extract_education(text: str, sections: dict):
    """Return (education_string, level_key, field) -- level/field may be None."""
    scope = sections.get("education") or text
    level = None
    for lvl, pattern in DEGREE_PATTERNS:
        if re.search(pattern, scope, flags=re.I):
            level = lvl
            break
    if level is None and scope is not text:
        for lvl, pattern in DEGREE_PATTERNS:  # degree might sit outside the education section
            if re.search(pattern, text, flags=re.I):
                level, scope = lvl, text
                break

    field = None
    if level:
        degree_re = dict(DEGREE_PATTERNS)[level]
        for line in scope.splitlines():
            if re.search(degree_re, line, flags=re.I):
                on_line = sorted((line.lower().find(f.lower()), -len(f), f) for f in KNOWN_FIELDS if mentions(line, f))
                if on_line:
                    field = on_line[0][2]
                    break
    hits = [] if field else sorted((scope.lower().find(f.lower()), f) for f in KNOWN_FIELDS if mentions(scope, f))
    if hits:
        # keep the most specific field (e.g. "Computer Science" over "Engineering") among the earliest mentions
        field = min(hits, key=lambda h: (h[0] // 80, -len(h[1])))[1]
    if field is None:
        m = re.search(r"(?:in|of)\s+([A-Z][A-Za-z&\- ]{2,40}?)(?=[,.;\n(|]|\s+(?:from|at|university|college|institute)\b)", scope)
        if m:
            field = m.group(1).strip()

    parts = [EDU_LABEL[level]] if level else []
    if field:
        parts.append(f"in {field}")
    return (" ".join(parts) if parts else ""), level, field


def extract_projects(sections: dict, limit: int = 6) -> list:
    body = sections.get("projects", "")
    if not body:
        return []
    has_markers = bool(re.search(r"^\s*(?:[-*–—]|\d+[.)])\s+", body, flags=re.M))
    raw_entries, current = [], []
    for line in body.splitlines():
        if not has_markers and line.strip() and re.match(r"^[A-Z][^:\n]{2,60}(?::|\s[–—-]\s)", line.strip()):
            if current:
                raw_entries.append(" ".join(current))
            current = [line.strip()]
        elif re.match(r"^\s*(?:[-*–—]|\d+[.)])\s+", line) or (not line.strip() and current):
            if current:
                raw_entries.append(" ".join(current))
            current = [] if not line.strip() else [re.sub(r"^\s*(?:[-*–—]|\d+[.)])\s+", "", line).strip()]
        elif line.strip():
            current.append(line.strip())
    if current:
        raw_entries.append(" ".join(current))

    projects = []
    for entry in raw_entries[:limit]:
        m = re.match(r"^(.{3,70}?)\s*(?::|\s[–—-]\s|\|)\s*(.+)$", entry)
        title, desc = (m.group(1), m.group(2)) if m else (entry[:60], entry)
        projects.append({"title": title.strip(), "description": desc.strip()})
    return projects


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------
def parse_resume(text: str, filename: str = "resume.txt", extra_terms=(), today: date = None) -> dict:
    """Parse resume text into the record format used by feature_engineering.compute_features."""
    warnings = []
    text = _clean(text)
    if len(text) < 150:
        warnings.append("Very little text was extracted. The file may be a scanned image (OCR is not supported).")

    sections = split_sections(text)
    skills = detect_skills(text, extra_terms)
    if not skills:
        warnings.append("No recognizable skills were detected.")

    exp_scope = sections.get("experience")
    if exp_scope is None:
        exp_scope = text.replace(sections.get("education", "\0"), "")  # keep education dates out
    from_dates = years_from_date_ranges(exp_scope, today)
    stated = years_stated(text)
    years = max(from_dates, stated)
    if years == 0:
        warnings.append("Could not determine years of experience; assuming 0.")
    method = "date ranges" if from_dates >= stated and from_dates > 0 else ("stated in text" if stated else "unknown")

    education, level, field = extract_education(text, sections)
    if level is None:
        warnings.append("No degree detected; education is scored as a bachelor's-equivalent default.")
        education = education or "Bachelor's (assumed)"
    if field is None:
        warnings.append("Field of study not detected; only degree level is compared.")

    projects = extract_projects(sections)
    if not projects:
        warnings.append("No 'Projects' section found; project relevance scores 0 for this resume.")

    return {
        "name": extract_name(text, filename),
        "email": extract_email(text),
        "skills": skills,
        "technologies": list(skills),  # a free-form resume doesn't separate "skills" from "tools"
        "experience_years": round(years, 1),
        "experience_method": method,
        "education": education,
        "education_level": level,
        "education_field": field,
        "projects": projects,
        "resume_text": text,
        "warnings": warnings,
    }
